"""XLSX fallback extraction follows the workbook, not worksheet filenames."""

import io
import zipfile

import pytest

from deeptutor.utils import document_extractor


def _workbook(*, metadata: bool = True, relationships: bool = True, target_mode: str = "") -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            "</Types>",
        )
        if metadata:
            archive.writestr(
                "xl/workbook.xml",
                '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<sheets><sheet name="Summary" sheetId="42" r:id="rSummary"/>'
                '<sheet name="Raw data" sheetId="7" r:id="rData"/></sheets></workbook>',
            )
        if relationships:
            archive.writestr(
                "xl/_rels/workbook.xml.rels",
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rData" Target="/xl/worksheets/sheet2.xml" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"/>'
                '<Relationship Id="rSummary" Target="worksheets/summary.xml" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
                f"{target_mode}/>"
                "</Relationships>",
            )
        for member, text in (
            ("sheet1.xml", "Unreferenced worksheet"),
            ("sheet2.xml", "Raw measurements"),
            ("summary.xml", "Study summary"),
        ):
            archive.writestr(
                f"xl/worksheets/{member}",
                '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                '<sheetData><row r="1"><c r="A1" t="inlineStr">'
                f"<is><t>{text}</t></is></c></row></sheetData></worksheet>",
            )
    return buffer.getvalue()


def test_fallback_resolves_sheet_names_and_order_through_relationships(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from openpyxl import load_workbook

    data = _workbook()
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Summary", "Raw data"]
        assert [sheet["A1"].value for sheet in workbook] == ["Study summary", "Raw measurements"]
    finally:
        workbook.close()
    monkeypatch.setattr(document_extractor, "load_workbook", None)

    text = document_extractor.extract_text_from_bytes("study.xlsx", data)

    assert text == (
        "--- Sheet: Summary ---\nStudy summary\n\n--- Sheet: Raw data ---\nRaw measurements"
    )


def test_fallback_skips_external_sheet_relationships(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(document_extractor, "load_workbook", None)

    text = document_extractor.extract_text_from_bytes(
        "study.xlsx", _workbook(target_mode='TargetMode="External"')
    )

    assert text == "--- Sheet: Raw data ---\nRaw measurements"


@pytest.mark.parametrize("missing", ["metadata", "relationships"])
def test_metadata_free_archives_keep_numeric_sheet_fallback(
    monkeypatch: pytest.MonkeyPatch, missing: str
) -> None:
    monkeypatch.setattr(document_extractor, "load_workbook", None)

    text = document_extractor.extract_text_from_bytes("study.xlsx", _workbook(**{missing: False}))

    assert text == (
        "--- Sheet: sheet1 ---\nUnreferenced worksheet\n\n--- Sheet: sheet2 ---\nRaw measurements"
    )
