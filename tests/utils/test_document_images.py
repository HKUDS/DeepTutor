"""Tests for deeptutor.utils.document_images — embedded-image extraction.

Fixtures generate real DOCX/PPTX archives on the fly (same pattern as
``test_document_extractor``) with a synthetic PNG large enough to clear the
minimum-size budget.
"""

from __future__ import annotations

import io
import random

from docx import Document as DocxDocument
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PptxInches
import pytest

from deeptutor.utils.document_images import (
    MAX_INLINE_IMAGE_EDGE,
    ImageBudget,
    build_marker,
    extract_docx_rich,
    extract_pdf_images,
    extract_pptx_rich,
    find_markers,
    fit_inline_image,
    image_index_from_name,
    reading_image_budget,
)


def _png_bytes(size: tuple[int, int] = (400, 300)) -> bytes:
    """A noisy PNG comfortably above the minimum-bytes budget."""
    rng = random.Random(11)
    image_bytes = bytes(rng.randrange(256) for _ in range(size[0] * size[1] * 3))
    from PIL import Image as PILImage

    image = PILImage.frombytes("RGB", size, image_bytes)
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture(scope="module")
def png() -> bytes:
    return _png_bytes()


def _docx_with_image(paragraphs: list[str], png_data: bytes, *, table: bool = False) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph(paragraphs[0])
    doc.add_picture(io.BytesIO(png_data), width=Inches(2))
    for paragraph in paragraphs[1:]:
        doc.add_paragraph(paragraph)
    if table:
        table_rows = doc.add_table(rows=1, cols=2)
        table_rows.cell(0, 0).text = "cell-a"
        table_rows.cell(0, 1).text = "cell-b"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_with_image(png_data: bytes) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(PptxInches(1), PptxInches(1), PptxInches(4), PptxInches(1))
    box.text_frame.text = "标题文本"
    slide.shapes.add_picture(io.BytesIO(png_data), PptxInches(1), PptxInches(2), PptxInches(3))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Markers
# ---------------------------------------------------------------------------


def test_marker_roundtrip() -> None:
    from deeptutor.utils.document_images import EmbeddedImage

    image = EmbeddedImage(name="image-03.png", mime_type="image/png", data=b"x")
    marker = build_marker(image)
    assert find_markers(f"前文\n{marker}\n后文") == [(3, "image-03.png")]


def test_image_index_from_name() -> None:
    assert image_index_from_name("image-01.png") == 1
    assert image_index_from_name("media/image-12.jpg") == 12
    assert image_index_from_name("nope.png") == 0


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def test_docx_rich_extracts_image_and_marker(png: bytes) -> None:
    data = _docx_with_image(["第一段", "第二段"], png)
    rich = extract_docx_rich(data)

    assert len(rich.collection.images) == 1
    image = rich.collection.images[0]
    assert image.mime_type == "image/png"
    assert image.data == png

    marker = build_marker(image)
    assert any(marker in paragraph for paragraph in rich.paragraphs)
    # Text order preserved: intro before the marker, outro after it.
    text = "\n".join(rich.paragraphs)
    assert text.index("第一段") < text.index(marker) < text.index("第二段")


def test_docx_rich_keeps_table_text(png: bytes) -> None:
    data = _docx_with_image(["开头"], png, table=True)
    rich = extract_docx_rich(data)
    joined = "\n".join(rich.paragraphs)
    assert "cell-a" in joined and "cell-b" in joined


def test_docx_rich_repeats_marker_for_reused_target(png: bytes) -> None:
    """The same image embedded twice is stored once but marked twice."""
    doc = DocxDocument()
    doc.add_picture(io.BytesIO(png), width=Inches(2))
    doc.add_paragraph("中间")
    doc.add_picture(io.BytesIO(png), width=Inches(2))
    buf = io.BytesIO()
    doc.save(buf)

    rich = extract_docx_rich(buf.getvalue())
    assert len(rich.collection.images) == 1
    markers = find_markers("\n".join(rich.paragraphs))
    assert len(markers) == 2


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


def test_pptx_rich_extracts_image_and_slide_text(png: bytes) -> None:
    rich = extract_pptx_rich(_pptx_with_image(png))

    assert len(rich.slides) == 1
    assert "标题文本" in rich.slides[0]
    assert len(rich.collection.images) == 1
    assert build_marker(rich.collection.images[0]) in rich.slides[0]


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def test_pdf_images_extracts_and_maps_pages(tmp_path) -> None:
    pymupdf = pytest.importorskip("pymupdf")
    from PIL import Image as PILImage

    png_data = _png_bytes((300, 200))
    pdf_path = tmp_path / "figures.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(pymupdf.Rect(72, 72, 300, 220), stream=png_data)
    doc.save(str(pdf_path))
    doc.close()
    del PILImage

    result = extract_pdf_images(pdf_path.read_bytes())
    assert result.collection.images, "the inserted figure should be extracted"
    assert result.page_map
    page_number, indices = result.page_map[0]
    assert page_number == 1
    assert result.collection.images[indices[0]].data


def test_pdf_images_deduplicates_repeated_xref(tmp_path) -> None:
    pytest.importorskip("pymupdf")
    png_data = _png_bytes((300, 200))
    pdf_path = tmp_path / "banner.pdf"
    doc = __import__("pymupdf").open()
    for _ in range(3):
        page = doc.new_page(width=612, height=792)
        page.insert_image(__import__("pymupdf").Rect(72, 72, 300, 160), stream=png_data)
    doc.save(str(pdf_path))
    doc.close()

    result = extract_pdf_images(pdf_path.read_bytes())
    assert len(result.collection.images) == 1
    assert result.page_map == ((1, (0,)), (2, (0,)), (3, (0,)))


def _png_for_index(index: int) -> bytes:
    """Distinct (content-wise) noisy PNG per index, all 64px+ and >min bytes."""
    return _png_bytes((120 + index, 120))


def _pdf_with_images(tmp_path, count: int) -> bytes:
    """One page carrying *count* distinct embedded images."""
    pymupdf = pytest.importorskip("pymupdf")
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    for index in range(count):
        rect = pymupdf.Rect(72, 72 + index * 40, 180, 108 + index * 40)
        page.insert_image(rect, stream=_png_for_index(index))
    pdf_path = tmp_path / f"many-{count}.pdf"
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path.read_bytes()


def test_pdf_images_respects_per_page_cap(tmp_path) -> None:
    """A page stops admitting at max_images_per_page; the rest are counted."""
    data = _pdf_with_images(tmp_path, 6)
    result = extract_pdf_images(data, budget=ImageBudget(max_images_per_page=4))

    assert len(result.collection.images) == 4
    assert result.collection.skipped_page_budget == 2
    assert result.collection.skipped_budget == 0
    note = result.collection.summary_note()
    assert "2 张超出单页上限未提取" in note


def test_pdf_images_honours_custom_max_images(tmp_path) -> None:
    data = _pdf_with_images(tmp_path, 6)
    result = extract_pdf_images(data, budget=ImageBudget(max_images=2))

    assert len(result.collection.images) == 2
    assert result.collection.skipped_budget == 4
    assert result.collection.skipped_page_budget == 0


def test_pdf_images_default_budget_caps_at_twelve(tmp_path) -> None:
    """No budget argument keeps the historical document-wide 12-image cap."""
    data = _pdf_with_images(tmp_path, 14)
    result = extract_pdf_images(data)

    assert len(result.collection.images) == 12
    assert result.collection.skipped_budget == 2
    assert result.collection.skipped_page_budget == 0


# ---------------------------------------------------------------------------
# Inline transport bound
# ---------------------------------------------------------------------------


def _noise_jpeg(size: tuple[int, int], quality: int = 95, seed: int = 7) -> bytes:
    """A large noisy JPEG: inside every byte budget, past the long-edge bound."""
    from PIL import Image as PILImage

    raw = random.Random(seed).randbytes(size[0] * size[1] * 3)
    image = PILImage.frombytes("RGB", size, raw)
    buf = io.BytesIO()
    image.save(buf, format="JPEG", quality=quality)
    return buf.getvalue()


def test_inline_budget_fits_a_page_sized_raster() -> None:
    """A scan past the long edge is fitted before admission, aspect preserved."""
    from PIL import Image as PILImage

    source = _noise_jpeg((2000, 1200))
    budget = ImageBudget()
    admitted = budget.try_admit(".jpg", source)

    assert admitted is not None
    assert admitted.name.endswith(".jpg")
    assert admitted.mime_type == "image/jpeg"
    assert len(admitted.data) < len(source)
    with PILImage.open(io.BytesIO(admitted.data)) as fitted:
        assert max(fitted.size) <= MAX_INLINE_IMAGE_EDGE
        assert fitted.size == (1600, 960)
    # The budget accounts the bytes that will actually ship, not the source.
    assert budget.total_bytes == len(admitted.data)


def test_reading_budget_keeps_source_resolution() -> None:
    """Images stored for the reader pane are written as-is, never refitted."""
    source = _noise_jpeg((2000, 1200))
    admitted = reading_image_budget().try_admit(".jpg", source)

    assert admitted is not None
    assert admitted.data == source
    assert admitted.mime_type == "image/jpeg"


def test_inline_budget_leaves_small_images_untouched() -> None:
    """Under the bound, the original bytes and format survive unchanged."""
    source = _png_bytes((400, 300))
    admitted = ImageBudget().try_admit(".png", source)

    assert admitted is not None
    assert admitted.data == source
    assert admitted.mime_type == "image/png"


def test_inline_budget_bounds_office_pictures_too() -> None:
    """DOCX/PPTX pictures reach the same admission seam as PDF page rasters."""
    from PIL import Image as PILImage

    rich = extract_docx_rich(_docx_with_image(["A picture follows"], _noise_jpeg((2000, 1200))))

    assert rich.collection.images, "the embedded picture should be extracted"
    with PILImage.open(io.BytesIO(rich.collection.images[0].data)) as fitted:
        assert max(fitted.size) <= MAX_INLINE_IMAGE_EDGE


def test_fit_returns_undecodable_bytes_unchanged() -> None:
    """A size optimisation must never be why an attachment becomes unreadable."""
    payload = b"\x89PNG\r\n\x1a\n" + b"not really a png" * 64

    assert fit_inline_image(".png", payload, max_edge=MAX_INLINE_IMAGE_EDGE) == (".png", payload)


def test_fit_composites_transparency_onto_white() -> None:
    """JPEG has no alpha channel; transparent areas must not turn black."""
    from PIL import Image as PILImage
    from PIL import ImageDraw as PILImageDraw

    raw = random.Random(5).randbytes(2000 * 1200 * 3)
    image = PILImage.frombytes("RGB", (2000, 1200), raw).convert("RGBA")
    image.putalpha(0)
    PILImageDraw.Draw(image).rectangle([0, 0, 300, 300], fill=(255, 0, 0, 255))
    buf = io.BytesIO()
    image.save(buf, format="PNG")

    ext, data = fit_inline_image(".png", buf.getvalue(), max_edge=MAX_INLINE_IMAGE_EDGE)

    assert ext == ".jpg"
    with PILImage.open(io.BytesIO(data)) as fitted:
        assert fitted.mode == "RGB"
        assert max(fitted.size) <= MAX_INLINE_IMAGE_EDGE
        # A transparent corner composites onto white, not onto black.
        assert fitted.getpixel((fitted.size[0] - 4, fitted.size[1] - 4)) == (255, 255, 255)
        opaque = fitted.getpixel((20, 20))
        assert opaque[0] > 200 and opaque[1] < 80 and opaque[2] < 80


def test_scanned_page_rides_bounded_but_reads_whole(tmp_path) -> None:
    """The reported case: a whole-page scan, inline vs on-disk resolution."""
    pymupdf = pytest.importorskip("pymupdf")
    from PIL import Image as PILImage

    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(pymupdf.Rect(0, 0, 612, 792), stream=_noise_jpeg((2000, 1200)))
    pdf_path = tmp_path / "scan.pdf"
    doc.save(str(pdf_path))
    doc.close()
    data = pdf_path.read_bytes()

    inline = extract_pdf_images(data).collection.images
    stored = extract_pdf_images(data, budget=reading_image_budget()).collection.images

    assert inline and stored
    with PILImage.open(io.BytesIO(inline[0].data)) as fitted:
        assert max(fitted.size) <= MAX_INLINE_IMAGE_EDGE
    with PILImage.open(io.BytesIO(stored[0].data)) as whole:
        assert max(whole.size) > MAX_INLINE_IMAGE_EDGE
