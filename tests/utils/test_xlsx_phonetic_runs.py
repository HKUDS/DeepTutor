"""Phonetic (furigana) runs are not cell text in the raw XLSX fallback."""

import io
import zipfile

from openpyxl import Workbook, load_workbook

from deeptutor.utils.document_extractor import _extract_xlsx_ooxml

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _workbook_with_cell(cell: str, shared_strings: str | None = None) -> bytes:
    """Save a one-sheet workbook with openpyxl, then swap in a raw A1 cell."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cities"
    sheet["A1"] = "placeholder"
    stream = io.BytesIO()
    workbook.save(stream)

    placeholder = '<c r="A1" t="inlineStr"><is><t>placeholder</t></is></c>'
    patched = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source,
        zipfile.ZipFile(patched, "w") as target,
    ):
        for item in source.infolist():
            data = source.read(item.filename).decode("utf-8")
            if item.filename == "xl/worksheets/sheet1.xml":
                assert placeholder in data
                data = data.replace(placeholder, cell)
            elif shared_strings is not None and item.filename == "[Content_Types].xml":
                data = data.replace(
                    "</Types>",
                    '<Override PartName="/xl/sharedStrings.xml" ContentType="application/'
                    'vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>',
                )
            elif shared_strings is not None and item.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(
                    "</Relationships>",
                    '<Relationship Id="rIdSst" Type="http://schemas.openxmlformats.org/'
                    'officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
                    "</Relationships>",
                )
            target.writestr(item, data)
        if shared_strings is not None:
            target.writestr("xl/sharedStrings.xml", shared_strings)
    return patched.getvalue()


def _openpyxl_text(data: bytes) -> str:
    workbook = load_workbook(io.BytesIO(data), data_only=True)
    try:
        return str(workbook.active["A1"].value)
    finally:
        workbook.close()


def test_shared_string_phonetic_run_is_ignored() -> None:
    data = _workbook_with_cell(
        '<c r="A1" t="s"><v>0</v></c>',
        f'<sst xmlns="{_MAIN_NS}" count="1" uniqueCount="1"><si><t>東京</t>'
        '<rPh sb="0" eb="2"><t>トウキョウ</t></rPh><phoneticPr fontId="1"/></si></sst>',
    )
    assert _openpyxl_text(data) == "東京"
    assert _extract_xlsx_ooxml(data, "cities.xlsx") == "--- Sheet: Cities ---\n東京"


def test_rich_text_shared_string_keeps_runs_and_drops_phonetic_runs() -> None:
    data = _workbook_with_cell(
        '<c r="A1" t="s"><v>0</v></c>',
        f'<sst xmlns="{_MAIN_NS}" count="1" uniqueCount="1"><si>'
        "<r><t>大阪</t></r><r><rPr><b/></rPr><t>府</t></r>"
        '<rPh sb="0" eb="2"><t>オオサカ</t></rPh><rPh sb="2" eb="3"><t>フ</t></rPh>'
        "</si></sst>",
    )
    assert _openpyxl_text(data) == "大阪府"
    assert _extract_xlsx_ooxml(data, "cities.xlsx") == "--- Sheet: Cities ---\n大阪府"


def test_inline_string_phonetic_run_is_ignored() -> None:
    data = _workbook_with_cell(
        '<c r="A1" t="inlineStr"><is><t>京都</t><rPh sb="0" eb="2"><t>キョウト</t></rPh></is></c>'
    )
    assert _openpyxl_text(data) == "京都"
    assert _extract_xlsx_ooxml(data, "cities.xlsx") == "--- Sheet: Cities ---\n京都"
