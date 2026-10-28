"""Check real core/Solve wheel installation and data-preserving removal."""

from __future__ import annotations

import argparse
import hashlib
import os
from pathlib import Path
import subprocess
import tempfile
import venv

ROOT = Path(__file__).resolve().parents[1]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--wheelhouse", required=True, type=Path)
    parser.add_argument(
        "--extended",
        action="store_true",
        help="Also check Parsing, Research and Calendar isolation",
    )
    args = parser.parse_args()
    wheelhouse = args.wheelhouse.resolve()
    with tempfile.TemporaryDirectory(prefix="deeptutor-module-check-") as directory:
        scratch = Path(directory)
        environment = scratch / "venv"
        venv.create(environment, with_pip=True)
        python = environment / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
        home = scratch / "home"
        env = {**os.environ, "DEEPTUTOR_HOME": str(home)}
        env.pop("PYTHONPATH", None)

        def run(*argv):
            subprocess.run([str(python), *map(str, argv)], cwd=scratch, env=env, check=True)

        core = list(wheelhouse.glob("deeptutor-*-py3-none-any.whl"))
        if len(core) != 1:
            raise ValueError("Wheelhouse must contain exactly one core version")
        run("-m", "pip", "install", core[0])

        def module(action, name="solve"):
            run("-m", "deeptutor_cli", "module", action, name, "--wheelhouse", wheelhouse)

        def smoke(capability="chat"):
            run(
                ROOT / "scripts/check_minimal_install.py",
                "--home",
                home,
                "--capability",
                capability,
            )

        smoke()
        run(
            "-c",
            """
from deeptutor.runtime.home import get_runtime_home
from docx import Document
from openpyxl import Workbook, load_workbook
root = get_runtime_home()
document = Document()
document.add_paragraph("Core document creation")
document.save(root / "core.docx")
assert Document(root / "core.docx").paragraphs[0].text == "Core document creation"
workbook = Workbook()
workbook.active["A1"] = "Core spreadsheet creation"
workbook.save(root / "core.xlsx")
assert load_workbook(root / "core.xlsx").active["A1"].value == "Core spreadsheet creation"
""",
        )
        module("install")
        smoke("deep_solve")

        def snapshot():
            return {
                p.relative_to(home): hashlib.sha256(p.read_bytes()).hexdigest()
                for p in home.rglob("*")
                if p.is_file()
                and p.name != "modules.json"
                and p.suffix in {".json", ".sqlite3", ".db", ".docx", ".xlsx"}
            }

        before = snapshot()
        assert before
        module("uninstall")
        assert before == snapshot(), "Uninstall changed application data"
        smoke()
        module("install")
        smoke("deep_solve")
        if args.extended:
            for name in ("parsing", "research", "calendar"):
                module("install", name)
            run(
                "-c",
                """
import asyncio
from datetime import datetime, timedelta, timezone
from deeptutor.modules.catalog import installed
assert not installed("knowledge") and not installed("mastery")
from deeptutor.runtime.home import get_runtime_home
from deeptutor.services.parsing import get_parse_service
source = get_runtime_home() / "isolated-parser.txt"
source.write_text("Parsing independently of knowledge.")
for _ in range(2):
    assert "independently" in get_parse_service().parse(source, engine="text_only").markdown
from deeptutor.agents.research.pipeline import ResearchPipeline
from deeptutor.tools.research_tools import _client
assert _client() is not None
pipeline = ResearchPipeline.__new__(ResearchPipeline)
pipeline.kb_name = None
asyncio.run(pipeline._prepare_pageindex_tools())
from deeptutor.services.calendar import layers
layers.owner_key = lambda: "chat:fixture"
start = datetime.now(timezone.utc)
assert not layers.read_layers("fixture", start, start + timedelta(days=7), "UTC").reviews
""",
            )
            for name in ("research", "calendar", "parsing"):
                preserved = snapshot()
                module("uninstall", name)
                assert preserved == snapshot(), f"Removing {name} changed application data"
                smoke("deep_solve")
            run(
                "-c",
                """
from deeptutor.runtime.home import get_runtime_home
from docx import Document
from openpyxl import load_workbook
root = get_runtime_home()
assert Document(root / "core.docx").paragraphs[0].text == "Core document creation"
assert load_workbook(root / "core.xlsx").active["A1"].value == "Core spreadsheet creation"
""",
            )
        run("-m", "pip", "check")
        print("Core/Solve install, run, replay, uninstall, retained data, reinstall: OK")


if __name__ == "__main__":
    main()
